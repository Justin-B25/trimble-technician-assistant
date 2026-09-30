"""Trimble internal DOCX styling — matches NPI Release Training Notes / Verso tab format."""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = 'Open Sans'
BLUE = RGBColor(0x00, 0x54, 0x8C)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB0, 0x00, 0x20)
WARN = RGBColor(0x8D, 0x6E, 0x00)
HEADER_FILL = '00548C'
ROW_ALT_FILL = 'F2F6FA'


def apply_document_defaults(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    rpr = style._element.rPr
    if rpr.rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    else:
        rFonts = rpr.rFonts
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    rFonts.set(qn('w:eastAsia'), FONT)


def set_run_font(run, size=11, bold=False, italic=False, color=None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    rpr = run._element.rPr
    rFonts = rpr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    rFonts.set(qn('w:eastAsia'), FONT)


def shade_cell(cell, hex_color: str) -> None:
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shd)


def add_title_block(
    doc: Document,
    title: str,
    subtitle: str = '',
    confidential: str = 'CONFIDENTIAL — Trimble internal use only. Do not share or distribute.',
) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    set_run_font(r, 18, True, color=BLUE)

    if subtitle:
        p2 = doc.add_paragraph()
        r2 = p2.add_run(subtitle)
        set_run_font(r2, 12, True, color=DARK)

    if confidential:
        p3 = doc.add_paragraph()
        r3 = p3.add_run(confidential)
        set_run_font(r3, 10, True, color=RED)


def add_meta_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    add_styled_table(doc, ('', ''), rows)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, 13, True, color=BLUE)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '10')
    bottom.set(qn('w:space'), '3')
    bottom.set(qn('w:color'), HEADER_FILL)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_subheading(doc: Document, text: str, size: int = 12) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size, True, color=DARK)


def add_body(doc: Document, text: str, *, bold=False, size=11, color=None, italic=False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size, bold, italic, color or DARK)


def add_bullet(doc: Document, text: str, *, size=11, color=None) -> None:
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size, color=color or DARK)


def add_label_block(doc: Document, label: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f'{label}: ')
    set_run_font(r, 11, True, color=DARK)
    if len(lines) == 1:
        r2 = p.add_run(lines[0])
        set_run_font(r2, 11, color=DARK)
    else:
        for line in lines:
            add_bullet(doc, line)


def add_warning_block(doc: Document, title: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    set_run_font(r, 11, True, color=WARN)
    for line in lines:
        add_bullet(doc, line)


def add_stop_block(doc: Document, title: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    set_run_font(r, 11, True, color=RED)
    for line in lines:
        add_bullet(doc, line)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=10) -> None:
    cell.text = ''
    r = cell.paragraphs[0].add_run(str(text))
    set_run_font(r, size, bold, color=color or DARK)


def add_styled_table(doc: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = 'Table Grid'
    for j, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], header, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
        shade_cell(table.rows[0].cells[j], HEADER_FILL)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            set_cell_text(table.rows[i].cells[j], val, bold=(j == 0), size=10)
            if i % 2 == 0:
                shade_cell(table.rows[i].cells[j], ROW_ALT_FILL)
    doc.add_paragraph()


def add_horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), HEADER_FILL)
    pBdr.append(bottom)
    pPr.append(pBdr)
