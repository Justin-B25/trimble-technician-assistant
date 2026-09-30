#!/usr/bin/env python3
"""Generate TMC Bench Crane Assembly Guide as a Word document from guide-export.json."""

from __future__ import annotations

import json
import subprocess
import sys
import urllib.error
import urllib.request
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from trimble_doc_style import (
    GRAY,
    RED,
    add_body,
    add_bullet,
    add_horizontal_rule,
    add_label_block,
    add_meta_table,
    add_section_heading,
    add_stop_block,
    add_styled_table,
    add_subheading,
    add_title_block,
    add_warning_block,
    apply_document_defaults,
    set_run_font,
)

ROOT = Path(__file__).resolve().parent.parent
JSON_PATH = ROOT / 'scripts' / 'guide-export.json'
EXPORT_SCRIPT = ROOT / 'scripts' / 'export-guide-json.py'
IMAGE_CACHE_DIR = ROOT / 'scripts' / 'image-cache'
OUTPUT_PATH = ROOT / 'TMC-Bench-Crane-Assembly-Guide.docx'
IMAGE_WIDTH = Inches(6.0)
USER_AGENT = 'TMC-Bench-Crane-DOCX-Export/1.0'

STATUS_LABELS = {
    'complete': 'Ready',
    'in-progress': 'In progress',
    'placeholder': 'Coming soon',
}


class ImageCache:
    """Download and cache Google Drive step photos for embedding in the DOCX."""

    def __init__(self, drive_images: dict[str, str], cache_dir: Path) -> None:
        self.drive_images = drive_images or {}
        self.cache_dir = cache_dir
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        self._resolved: dict[str, Path | None] = {}
        self.stats = {'embedded': 0, 'missing': 0, 'failed': 0}

    def file_id_for_key(self, key: str) -> str:
        if not key:
            return ''
        return str(self.drive_images.get(key, '') or '').strip()

    def download_urls(self, file_id: str) -> list[str]:
        return [
            f'https://drive.google.com/thumbnail?id={file_id}&sz=w1600',
            f'https://lh3.googleusercontent.com/d/{file_id}=w1600',
            f'https://drive.google.com/uc?export=download&id={file_id}',
        ]

    def _guess_extension(self, content_type: str, data: bytes) -> str:
        ct = (content_type or '').lower()
        if 'jpeg' in ct or 'jpg' in ct:
            return '.jpg'
        if 'webp' in ct:
            return '.webp'
        if 'gif' in ct:
            return '.gif'
        if data[:8] == b'\x89PNG\r\n\x1a\n':
            return '.png'
        if data[:2] == b'\xff\xd8':
            return '.jpg'
        if data[:4] == b'RIFF' and data[8:12] == b'WEBP':
            return '.webp'
        return '.png'

    def _fetch(self, file_id: str) -> Path | None:
        for url in self.download_urls(file_id):
            try:
                req = urllib.request.Request(url, headers={'User-Agent': USER_AGENT})
                with urllib.request.urlopen(req, timeout=45) as response:
                    data = response.read()
                if len(data) < 1024:
                    continue
                if b'<html' in data[:256].lower():
                    continue
                ext = self._guess_extension(response.headers.get('Content-Type', ''), data)
                path = self.cache_dir / f'{file_id}{ext}'
                path.write_bytes(data)
                return path
            except (urllib.error.URLError, TimeoutError, OSError):
                continue
        return None

    def get_path(self, key: str) -> Path | None:
        file_id = self.file_id_for_key(key)
        if not file_id:
            return None
        if file_id in self._resolved:
            return self._resolved[file_id]

        for path in self.cache_dir.glob(f'{file_id}.*'):
            if path.is_file() and path.stat().st_size > 1024:
                self._resolved[file_id] = path
                return path

        path = self._fetch(file_id)
        self._resolved[file_id] = path
        return path


def drive_image_url(image_id: str) -> str:
    if not image_id:
        return ''
    return f'https://drive.google.com/file/d/{image_id}/view'


def add_embedded_image(
    doc: Document,
    cache: ImageCache,
    image_key: str,
    caption: str = '',
    placeholder: str = '',
) -> None:
    p = doc.add_paragraph()
    r = p.add_run('Reference photo')
    set_run_font(r, 11, True)

    path = cache.get_path(image_key)
    file_id = cache.file_id_for_key(image_key)

    if path:
        pic_para = doc.add_paragraph()
        pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_para.add_run().add_picture(str(path), width=IMAGE_WIDTH)
        cache.stats['embedded'] += 1
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(caption)
            set_run_font(r, 9, italic=True, color=GRAY)
        if file_id:
            link = doc.add_paragraph()
            link.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = link.add_run(drive_image_url(file_id))
            set_run_font(r, 8, color=GRAY)
        return

    if file_id:
        cache.stats['failed'] += 1
        add_body(
            doc,
            f'Photo could not be downloaded for embedding. Open online: {drive_image_url(file_id)}',
            size=10,
            color=GRAY,
        )
    elif placeholder:
        cache.stats['missing'] += 1
        add_body(doc, placeholder, size=10, italic=True, color=GRAY)
    else:
        cache.stats['missing'] += 1
        add_body(doc, 'Reference photo not yet available for this step.', size=10, italic=True, color=GRAY)


def refresh_json() -> None:
    if not EXPORT_SCRIPT.exists():
        raise FileNotFoundError(f'Missing export script: {EXPORT_SCRIPT}')
    subprocess.run([sys.executable, str(EXPORT_SCRIPT)], check=True, cwd=ROOT)


def load_data() -> dict:
    refresh_json()
    with JSON_PATH.open(encoding='utf-8') as f:
        return json.load(f)


def format_hw_line(entry: dict, catalog: dict) -> str:
    ref = entry.get('ref')
    qty = entry.get('qty')
    cat = catalog.get(ref, {}) if ref else {}
    name = cat.get('formalName') or entry.get('part') or ref or 'Hardware'
    pn = cat.get('partNumber')
    bits = [name]
    if pn:
        bits.append(f'P/N {pn}')
    if qty is not None and qty != '':
        bits.append(f'qty {qty}')
    note = entry.get('note') or cat.get('notes')
    if note:
        bits.append(note)
    return ' — '.join(bits)


def add_print_parts_table(doc: Document, parts: list[dict]) -> None:
    if not parts:
        return
    rows = []
    for part in parts:
        notes = ' · '.join(n for n in [part.get('folder'), part.get('note')] if n)
        rows.append(
            (
                str(part.get('name', '')),
                str(part.get('file', '')),
                str(part.get('qty', '')),
                notes,
            )
        )
    add_styled_table(doc, ('Part name', 'STL file', 'Qty', 'Notes'), rows)


def add_hardware_table(doc: Document, items: list[dict], catalog: dict) -> None:
    if not items:
        return
    rows = [(str(item.get('qty', '')), format_hw_line(item, catalog)) for item in items]
    add_styled_table(doc, ('Qty', 'Part / specification'), rows)


def render_step(doc: Document, step: dict, step_num: int, catalog: dict, cache: ImageCache) -> None:
    title = step.get('title', 'Step')
    req = ' (REQUIRED)' if step.get('required') else ''
    add_subheading(doc, f'Step {step_num}: {title}{req}', size=11)

    if step.get('summary'):
        add_body(doc, step['summary'])

    if step.get('configs'):
        add_label_block(doc, 'Applies to configuration', step['configs'])

    if step.get('warnings'):
        add_warning_block(doc, 'Warnings', step['warnings'])

    if step.get('orientationNotes'):
        add_label_block(doc, 'Procedure', step['orientationNotes'])

    if step.get('legend'):
        add_label_block(
            doc,
            'Reference legend',
            [f"{item.get('color', '').title()}: {item.get('label', '')}" for item in step['legend']],
        )

    if step.get('printedParts'):
        add_body(doc, 'Printed parts', bold=True, size=10)
        add_print_parts_table(doc, step['printedParts'])

    hw = step.get('hardware') or []
    if step.get('hardwareByVruConfig'):
        add_body(doc, 'Hardware (VRU configuration dependent)', bold=True, size=10)
        for cfg, items in step['hardwareByVruConfig'].items():
            add_body(doc, cfg, bold=True, size=10)
            add_hardware_table(doc, items, catalog)
    elif step.get('hardwarePerSide'):
        add_body(doc, 'Hardware per side', bold=True, size=10)
        add_hardware_table(doc, step['hardwarePerSide'], catalog)
    elif step.get('hardwareTotals'):
        add_body(doc, 'Hardware totals', bold=True, size=10)
        for group in step['hardwareTotals'].values():
            add_body(doc, group.get('label', 'Items'), bold=True, size=10)
            items = group.get('items', [])
            if items and isinstance(items[0], dict):
                add_hardware_table(doc, items, catalog)
            else:
                for key, val in group.items():
                    if key != 'label':
                        add_bullet(doc, f'{key}: {val}')
    elif hw:
        add_body(doc, 'Hardware', bold=True, size=10)
        add_hardware_table(doc, hw, catalog)

    if step.get('checks'):
        add_label_block(doc, 'Checks before continuing', step['checks'])

    addendum = step.get('optionalAddendum')
    if addendum:
        add_subheading(doc, addendum.get('title', 'Optional'), size=11)
        if addendum.get('summary'):
            add_body(doc, addendum['summary'])
        if addendum.get('orientationNotes'):
            add_label_block(doc, 'Procedure', addendum['orientationNotes'])
        if addendum.get('hardware'):
            add_hardware_table(doc, addendum['hardware'], catalog)
        if addendum.get('checks'):
            add_label_block(doc, 'Checks', addendum['checks'])

    if step.get('imageDriveId') or step.get('imagePlaceholder'):
        add_embedded_image(
            doc,
            cache,
            step.get('imageDriveId') or '',
            caption=step.get('imageAlt', ''),
            placeholder=step.get('imagePlaceholder', ''),
        )

    if step.get('imageNote'):
        add_body(doc, step['imageNote'], size=10, italic=True, color=GRAY)


def render_section(doc: Document, section_id: str, section: dict, catalog: dict, cache: ImageCache) -> None:
    title = section.get('title', section_id)
    status = STATUS_LABELS.get(section.get('status', ''), section.get('status', ''))
    add_section_heading(doc, title)
    add_body(doc, f'Segment status: {status}', size=10, color=GRAY)
    if section.get('driveFolder'):
        add_body(doc, f'STL library folder: {section["driveFolder"]}', size=10, color=GRAY)

    if section.get('assemblyNote'):
        add_body(doc, section['assemblyNote'])

    banner = section.get('criticalBanner')
    if banner:
        add_stop_block(doc, banner.get('title', 'Important'), banner.get('lines', []))

    if section.get('prerequisitesBeforeExtrusions'):
        add_label_block(doc, 'Prerequisites before extrusions', section['prerequisitesBeforeExtrusions'])

    printing = section.get('printingGuidance')
    if printing:
        add_label_block(doc, printing.get('title', 'Print tips'), printing.get('lines', []))

    if section.get('configuration'):
        cfg = section['configuration']
        add_body(doc, cfg.get('prompt', 'Configuration'), bold=True)
        for opt in cfg.get('options', []):
            default = ' (default)' if opt.get('default') else ''
            add_bullet(doc, f"{opt.get('label', '')}{default}: {opt.get('description', '')}")

    if section.get('vruConfiguration'):
        vru = section['vruConfiguration']
        add_body(doc, vru.get('prompt', 'VRU configuration'), bold=True)
        for opt in vru.get('options', []):
            default = ' (default)' if opt.get('default') else ''
            add_bullet(doc, f"{opt.get('label', '')}{default}: {opt.get('description', '')}")

    if section.get('assemblyMethod'):
        am = section['assemblyMethod']
        add_label_block(doc, am.get('title', 'Assembly method'), [am.get('summary', '')])
        if am.get('sides'):
            for line in am['sides']:
                add_bullet(doc, line)

    if section.get('chapterHardware'):
        add_body(doc, 'Segment hardware summary', bold=True, size=10)
        add_hardware_table(doc, section['chapterHardware'], catalog)

    print_parts = section.get('printParts') or section.get('sharedPrintedParts')
    if print_parts:
        add_body(doc, 'Print checklist', bold=True, size=10)
        add_print_parts_table(doc, print_parts)

    config_parts = section.get('configPrintedParts')
    if config_parts:
        for cfg_id, block in config_parts.items():
            add_body(doc, f'Additional prints — {cfg_id}', bold=True, size=10)
            add_print_parts_table(doc, block.get('include', []))
            if block.get('excludeNote'):
                add_body(doc, block['excludeNote'], size=10, italic=True, color=GRAY)

    steps = section.get('steps') or []
    if not steps and section.get('placeholderMessage'):
        add_body(doc, section['placeholderMessage'], italic=True, color=GRAY)
    for i, step in enumerate(steps, start=1):
        render_step(doc, step, i, catalog, cache)

    add_horizontal_rule(doc)


def build_document(data: dict) -> tuple[Document, ImageCache]:
    guide = data['GUIDE']
    catalog = data['HARDWARE_CATALOG']
    master_kit = data['MASTER_HARDWARE_KIT']
    order = data['ASSEMBLY_ORDER']
    cache = ImageCache(data.get('CRANE_DRIVE_IMAGES', {}), IMAGE_CACHE_DIR)

    doc = Document()
    apply_document_defaults(doc)

    add_title_block(
        doc,
        'TMC Bench Crane — Assembly Guide',
        guide.get('subtitle', 'Assembly procedure and parts list'),
    )
    add_meta_table(
        doc,
        [
            ('Product', 'TMC Bench Wire Crane'),
            ('Document type', 'Assembly guide export'),
            ('Generated', date.today().isoformat()),
            ('Interactive source', 'TMC Bench Crane guide-data.js'),
        ],
    )
    add_body(
        doc,
        'This document was exported from the interactive assembly guide. '
        'Key step reference photos are embedded for offline use.',
    )
    if guide.get('viewer3d', {}).get('url'):
        add_body(doc, f"360° 3D model (Fusion): {guide['viewer3d']['url']}", size=10, color=GRAY)

    doc.add_page_break()

    overview = guide['sections']['overview']
    add_section_heading(doc, 'Overview')
    intro = overview.get('intro', {})
    add_subheading(doc, intro.get('title', 'How this guide is organized'))
    for paragraph in intro.get('paragraphs', []):
        add_bullet(doc, paragraph)

    tools = overview.get('requiredTools', {})
    add_section_heading(doc, tools.get('title', 'Required tools'))
    if tools.get('intro'):
        add_body(doc, tools['intro'])
    for tool in tools.get('items', []):
        add_body(doc, tool.get('name', 'Tool'), bold=True)
        if tool.get('preferred'):
            add_body(doc, tool['preferred'], size=10, color=GRAY)
        if tool.get('usedFor'):
            add_bullet(doc, f"Used for: {tool['usedFor']}")
        if tool.get('detail'):
            add_bullet(doc, tool['detail'])
        for w in tool.get('warnings', []):
            add_bullet(doc, f'Note: {w}')
        if tool.get('imageDriveId'):
            add_embedded_image(doc, cache, tool['imageDriveId'], caption=tool.get('imageAlt', ''))

    mpl = overview.get('masterPartsList', {})
    add_section_heading(doc, mpl.get('title', 'Master hardware kit'))
    if mpl.get('incompleteNotice'):
        add_warning_block(doc, 'Master list status', [mpl['incompleteNotice']])
    add_body(doc, guide.get('hardwareKitIntro', ''))
    add_hardware_table(
        doc,
        [{'ref': item['ref'], 'qty': '', 'note': item.get('note', '')} for item in master_kit],
        catalog,
    )

    add_section_heading(doc, 'Hardware catalog reference')
    catalog_rows = []
    for item in catalog.values():
        spec_bits = [
            item.get('partNumber') and f"P/N {item['partNumber']}",
            item.get('thread'),
            item.get('length'),
            item.get('installHole') and f"hole {item['installHole']}",
        ]
        catalog_rows.append(
            (
                item.get('formalName', ''),
                ' · '.join(b for b in spec_bits if b),
                item.get('usedIn', ''),
            )
        )
    add_styled_table(doc, ('Part', 'P/N / spec', 'Used in'), catalog_rows)

    doc.add_page_break()

    add_section_heading(doc, 'Assembly segments (in order)')
    add_body(doc, overview.get('segmentsIntro', ''))
    segment_rows = [
        (str(i), f"{guide['sections'][seg_id].get('title', seg_id)} — {STATUS_LABELS.get(guide['sections'][seg_id].get('status', ''), '')}")
        for i, seg_id in enumerate(order, start=1)
    ]
    add_styled_table(doc, ('#', 'Segment'), segment_rows)

    doc.add_page_break()

    for seg_id in order:
        render_section(doc, seg_id, guide['sections'][seg_id], catalog, cache)
        doc.add_page_break()

    add_section_heading(doc, 'Document control')
    add_body(
        doc,
        'Segments marked “Coming soon” or “In progress” reflect the current state of the interactive guide. '
        'Regenerate this document after guide updates to distribute the latest instructions.',
    )
    add_body(
        doc,
        'Distribution: Trimble authorized personnel only. Do not post publicly or share outside Trimble without approval.',
    )
    add_body(doc, 'CONFIDENTIAL — Internal assembly documentation only.', bold=True, size=9, color=RED)

    return doc, cache


def main() -> int:
    try:
        data = load_data()
    except subprocess.CalledProcessError:
        print('Failed to export JSON from guide-data.js', file=sys.stderr)
        return 1
    except FileNotFoundError as exc:
        print(exc, file=sys.stderr)
        return 1

    doc, cache = build_document(data)
    doc.save(OUTPUT_PATH)
    print(f'Wrote {OUTPUT_PATH}')
    print(
        f"Photos embedded: {cache.stats['embedded']} · "
        f"missing ID: {cache.stats['missing']} · "
        f"download failed: {cache.stats['failed']}"
    )
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
